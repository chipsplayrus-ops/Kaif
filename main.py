import streamlit as st
import io
import asyncio
from Test3 import HHParser
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd

st.set_page_config(page_title="HH Search", page_icon="⚡", layout="wide")

st.title("⚡ Поиск вакансий HH.ru")

# Инициализируем хранилище данных в сессии
if "raw_results" not in st.session_state:
    st.session_state.raw_results = None

with st.sidebar:
    st.header("🔑 Доступ")
    token = st.text_input("Access Token:", type="password",
                          value="APPLTRTV1CFK3PLE95GC1DUFUFDGDLQHQQ0948NE05AIOLHUN77LMOVUA7CUOT23")

c1, c2 = st.columns([3, 1])
with c1:
    search_query = st.text_input("Какую работу ищем?", value="Python Developer")
with c2:
    num_pages = st.number_input("Глубина (стр):", 1, 10, 1)


def create_excel(df):
    output = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Вакансии"

    # Стилизация
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    stripe_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    border = Border(left=Side(style='thin'), right=Side(style='thin'),
                    top=Side(style='thin'), bottom=Side(style='thin'))

    # Заголовки
    headers = list(df.columns)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    # Заполнение строками данных
    for r_idx, row in enumerate(df.values, 2):
        ws.append(list(row))
        for c_idx, cell in enumerate(ws[r_idx], 1):
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if r_idx % 2 == 0:
                cell.fill = stripe_fill
            if headers[c_idx - 1] == "Ссылка":
                cell.font = Font(color="0563C1", underline="single")
                cell.hyperlink = cell.value

    # Корректная ширина колонок
    column_widths = [25, 20, 15, 15, 10, 15, 15, 50, 30]
    for i, width in enumerate(column_widths, 1):
        if i <= len(headers):
            ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A2"
    wb.save(output)
    return output.getvalue()


# Кнопка запуска поиска
if st.button("🚀 Найти вакансии", type="primary", use_container_width=True):
    if not token or not search_query:
        st.error("Заполните поля поиска!")
    else:
        try:
            parser = HHParser(token)

            with st.spinner("⚡ Получение данных..."):
                loop = asyncio.new_event_loop()
                asyncio.set_event_loop(loop)
                results = loop.run_until_complete(parser.get_vacancies(search_query, pages=num_pages))

            if isinstance(results, str):
                st.error(results)
                st.session_state.raw_results = None
            elif results:
                st.session_state.raw_results = results
            else:
                st.warning("Ничего не найдено. Попробуйте другой запрос.")
                st.session_state.raw_results = None

        except Exception as e:
            st.error(f"Ошибка приложения: {e}")

if st.session_state.raw_results:
    df = pd.DataFrame(st.session_state.raw_results)

    df['Зарплата ОТ'] = pd.to_numeric(df['Зарплата ОТ'], errors='coerce')
    df['Зарплата ДО'] = pd.to_numeric(df['Зарплата ДО'], errors='coerce')

    st.markdown("---")
    st.subheader("⚙️ Фильтрация результатов")

    # 1. Фильтр по городу
    city_input = st.text_input("🏙️ Город:", value="")
    if city_input:
        df = df[df['Город'].str.contains(city_input, case=False, na=False)]

    # 2. Фильтры Зарплаты и Опыта
    filter_c1, filter_c2 = st.columns(2)

    with filter_c1:
        min_salary = st.number_input("💵 Минимальная зарплата (ОТ):", min_value=0, value=0, step=10000)
        if min_salary > 0:
            df = df[df['Зарплата ОТ'].fillna(0) >= min_salary]

    with filter_c2:
        # Собираем уникальные варианты опыта из загруженных данных
        available_exp = list(df['Опыт'].dropna().unique())
        selected_exp = st.multiselect("💼 Требуемый опыт:", options=available_exp, default=available_exp)
        if selected_exp:
            df = df[df['Опыт'].isin(selected_exp)]

    st.info(f"Отображено вакансий: {len(df)}")

    st.subheader("📊 Предпросмотр данных")
    st.dataframe(df, use_container_width=True, hide_index=True, height=500)

    st.divider()
    excel_file = create_excel(df)
    st.download_button(
        label="📥 Скачать отфильтрованный Excel отчет",
        data=excel_file,
        file_name=f"HH_Report_{search_query}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True
    )